"""文档解析：支持PDF/PPT/Word，扫描版PDF自动走OCR，老格式(.ppt/.doc)自动转换。"""
import os
import shutil
import subprocess
import tempfile
import fitz  # PyMuPDF
from pptx import Presentation
from docx import Document


def _find_soffice() -> str:
    """查找LibreOffice可执行文件，找不到返回空串。"""
    exe = shutil.which("soffice")
    if exe:
        return exe
    for candidate in (
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        "/usr/bin/libreoffice",
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    ):
        if os.path.exists(candidate):
            return candidate
    return ""


def _convert_office(path: str, target_ext: str) -> str:
    """用LibreOffice把老格式(.ppt/.doc)转换成.pptx/.docx，返回新文件路径。"""
    soffice = _find_soffice()
    if not soffice:
        raise RuntimeError(
            f"文件{os.path.basename(path)}是老格式，需要LibreOffice转换，但未检测到，请另存为新格式后重试。"
        )
    out_dir = tempfile.mkdtemp(prefix="mg_office_")
    cmd = [soffice, "--headless", "--convert-to", target_ext,
           "--outdir", out_dir, path]
    proc = subprocess.run(cmd, capture_output=True, timeout=120)
    base = os.path.splitext(os.path.basename(path))[0]
    out_path = os.path.join(out_dir, base + "." + target_ext)
    if not os.path.exists(out_path):
        raise RuntimeError(
            f"老格式转换失败：{proc.stderr.decode('utf-8', 'ignore')[:200]}"
        )
    return out_path


def parse_pdf(path: str) -> str:
    """提取PDF文本，文本过少时判定为扫描件并回退OCR。"""
    parts = []
    doc = fitz.open(path)
    for page in doc:
        parts.append(page.get_text())
    doc.close()
    text = "\n".join(parts).strip()
    if len(text) < 100:
        text = _ocr_pdf(path)
    return text


def _ocr_pdf(path: str) -> str:
    """扫描版PDF走RapidOCR，逐页转图片后识别。"""
    from rapidocr_onnxruntime import RapidOCR
    ocr = RapidOCR()
    doc = fitz.open(path)
    parts = []
    for page in doc:
        pix = page.get_pixmap(dpi=200)
        img_path = path + ".tmp.png"
        pix.save(img_path)
        result, _ = ocr(img_path)
        if result:
            parts.append("\n".join([line[1] for line in result]))
        try:
            os.remove(img_path)
        except OSError:
            pass
    doc.close()
    return "\n".join(parts)


def parse_pptx(path: str) -> str:
    """提取PPT每一页的文字。"""
    prs = Presentation(path)
    parts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in para.runs).strip()
                    if line:
                        parts.append(line)
    return "\n".join(parts)


def parse_docx(path: str) -> str:
    """提取Word段落文本。"""
    doc = Document(path)
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def parse_file(path: str) -> str:
    """根据扩展名分发到对应解析器，老格式(.ppt/.doc)先转换。"""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return parse_pdf(path)
    if ext == ".pptx":
        return parse_pptx(path)
    if ext == ".docx":
        return parse_docx(path)
    if ext == ".ppt":
        converted = _convert_office(path, "pptx")
        try:
            return parse_pptx(converted)
        finally:
            shutil.rmtree(os.path.dirname(converted), ignore_errors=True)
    if ext == ".doc":
        converted = _convert_office(path, "docx")
        try:
            return parse_docx(converted)
        finally:
            shutil.rmtree(os.path.dirname(converted), ignore_errors=True)
    raise ValueError(f"暂不支持的文件格式: {ext}")
